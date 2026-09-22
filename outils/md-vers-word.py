"""
Convertit une ou plusieurs fiches .md en .docx dans VERSION-WORD/ (même arborescence),
avec la mise en page des autres documents Word de la valise.

    python outils/md-vers-word.py 16-CADRE-DE-REFERENCE/Bibliographie-de-reference.md [...]
    python outils/md-vers-word.py --dossier 16-CADRE-DE-REFERENCE

Nécessite : pip install python-docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Emu

ROOT = Path(__file__).resolve().parent.parent
BLUE = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x55, 0x55, 0x55)

INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))')


def add_inline(par, text, bold=False, italic=False, color=None, size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        b, i, mono = bold, italic, False
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            part, b = part[2:-2], True
        elif part.startswith('`') and part.endswith('`'):
            part, mono = part[1:-1], True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            part, i = part[1:-1], True
        run = par.add_run(part)
        run.bold = b or None
        run.italic = i or None
        if mono:
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        if color is not None:
            run.font.color.rgb = color
        if size is not None:
            run.font.size = size


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def split_row(line):
    line = line.strip().strip('|')
    return [c.strip() for c in re.split(r'(?<!\\)\|', line)]


def convert(md_path: Path):
    lines = md_path.read_text(encoding='utf-8').splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Emu(731520)
    sec.top_margin = sec.bottom_margin = Emu(640080)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Emu(133350)

    i, first_h1 = 0, True
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith('```'):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            r = p.add_run('\n'.join(buf))
            r.font.name = 'Consolas'
            r.font.size = Pt(8.5)
            continue
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            lvl, txt = len(m.group(1)), m.group(2)
            if lvl == 1 and first_h1:
                first_h1 = False
                p = doc.add_paragraph()
                add_inline(p, txt, bold=True, color=BLUE, size=Emu(228600))
            else:
                doc.add_heading(re.sub(r'\*\*', '', txt), level=min(max(lvl, 2), 4))
            i += 1
            continue
        if re.match(r'^([-*_])(\s*\1){2,}$', s):
            doc.add_paragraph()
            i += 1
            continue
        if s.startswith('>'):
            while i < len(lines) and lines[i].strip().startswith('>'):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Emu(182880)
                add_inline(p, re.sub(r'^>\s?', '', lines[i].strip()), italic=True, color=GREY)
                i += 1
            continue
        if '|' in s and i + 1 < len(lines) and re.match(r'^\|?\s*:?-{2,}', lines[i + 1].strip()):
            head = split_row(s)
            i += 2
            rows = []
            while i < len(lines) and '|' in lines[i] and lines[i].strip():
                rows.append(split_row(lines[i]))
                i += 1
            t = doc.add_table(rows=1 + len(rows), cols=len(head))
            t.style = 'Table Grid'
            for k, h in enumerate(head):
                c = t.rows[0].cells[k]
                shade(c, '1F4E79')
                add_inline(c.paragraphs[0], h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=Pt(9.5))
            for r_i, row in enumerate(rows, start=1):
                for k in range(len(head)):
                    c = t.rows[r_i].cells[k]
                    add_inline(c.paragraphs[0], row[k] if k < len(row) else '', size=Pt(9.5))
                    if r_i % 2 == 0:
                        shade(c, 'EEF3F8')
            doc.add_paragraph()
            continue
        m = re.match(r'^(\s*)([-*+]|\d+[.)])\s+(.*)$', line)
        if m:
            ordered = m.group(2)[0].isdigit()
            nested = len(m.group(1).replace('\t', '    ')) >= 2
            style = ('List Number' if ordered else 'List Bullet') + (' 2' if nested else '')
            p = doc.add_paragraph(style=style)
            add_inline(p, re.sub(r'^\[[ xX]\]\s+', '☐ ', m.group(3)))
            i += 1
            continue
        p = doc.add_paragraph()
        buf = []
        while i < len(lines) and lines[i].strip() and not re.match(r'^(#|>|```|\s*([-*+]|\d+[.)])\s)', lines[i].strip()):
            if '|' in lines[i] and i + 1 < len(lines) and re.match(r'^\|?\s*:?-{2,}', lines[i + 1].strip()):
                break
            buf.append(lines[i].strip())
            i += 1
        for k, b in enumerate(buf):
            if k:
                p.add_run().add_break()
            add_inline(p, b)

    out = ROOT / 'VERSION-WORD' / md_path.relative_to(ROOT).with_suffix('.docx')
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print('OK :', out.relative_to(ROOT).as_posix())


if __name__ == '__main__':
    args = sys.argv[1:]
    paths = []
    if args[:1] == ['--dossier']:
        paths = sorted((ROOT / args[1]).glob('*.md'))
    else:
        paths = [ROOT / a for a in args]
    if not paths:
        print(__doc__)
        sys.exit(1)
    for p in paths:
        convert(p.resolve())
